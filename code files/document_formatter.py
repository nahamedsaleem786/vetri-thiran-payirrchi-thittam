from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parents[1]


def _set_run_font(run, size=11, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold


def _set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    _set_run_font(run, size=10, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_footer(document):
    for section in document.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Generated with LegalEase • AI-assisted draft")
        _set_run_font(run, size=8)


def _looks_like_heading(line: str) -> bool:
    if line.isupper() and len(line) <= 140:
        return True
    first = line.split(".", 1)[0]
    return first.isdigit() and len(first) <= 3


def format_docx(text: str, doc_type: str, terms: str = "", logo_path: str | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    logo = Path(logo_path) if logo_path else BASE_DIR / "assets" / "logo.png"
    if logo.exists():
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(logo), width=Inches(1.1))
        except Exception:
            pass

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    _set_run_font(run, size=16, bold=True)

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            doc.add_paragraph()
            continue

        p = doc.add_paragraph()
        run = p.add_run(line)
        _set_run_font(run, size=11, bold=_looks_like_heading(line))

    clean_terms = [t.strip() for t in terms.split(";") if t.strip()]
    if clean_terms:
        heading = doc.add_paragraph()
        r = heading.add_run("KEY TERMS")
        _set_run_font(r, size=11, bold=True)

        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        _set_cell_text(table.rows[0].cells[0], "No.", True)
        _set_cell_text(table.rows[0].cells[1], "Term / Condition", True)

        for i, term in enumerate(clean_terms, 1):
            cells = table.add_row().cells
            _set_cell_text(cells[0], str(i))
            _set_cell_text(cells[1], term)

    _add_footer(doc)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
